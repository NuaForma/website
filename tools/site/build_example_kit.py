"""Reproducibly author the downloadable website template and complete request envelopes."""
import copy
import argparse
import io
import json
import pathlib
import zipfile
from build_client_collections import build as build_client_collections
from package_complete_kit import package as package_complete_kit

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw
from lxml import etree

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = ROOT / "site/downloads"
FONT = "Liberation Serif"


def element(parent, tag, attributes):
    item = OxmlElement(tag)
    for name, value in attributes.items():
        item.set(qn(name), str(value))
    parent.append(item)
    return item


def normalize_zip(file):
    with zipfile.ZipFile(file) as source:
        contents = {item.filename: source.read(item) for item in source.infolist()}
    # Default numbering and stylesWithEffects also carry unused theme references.
    for name, data in list(contents.items()):
        if name.startswith("word/") and name.endswith(".xml"):
            root = etree.fromstring(data)
            # Author an opaque, effect-free theme as well as opaque visible images.
            for node in list(root.iter()):
                local = etree.QName(node).localname.lower()
                if local in ("effectlst", "effectdag"):
                    for child in list(node):
                        node.remove(child)
                elif local == "alpha":
                    node.set("val", "100000")
            fonts = list(root.iter(qn("w:rFonts")))
            for declaration in fonts:
                for key in list(declaration.attrib):
                    if key.lower().endswith("theme"):
                        del declaration.attrib[key]
                for family in ("ascii", "hAnsi", "eastAsia", "cs"):
                    declaration.set(qn("w:" + family), FONT)
            contents[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(contents.items()):
            item = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(item, data)
    file.write_bytes(output.getvalue())


def build_template(include_content_image=False):
    doc = Document()
    # The python-docx starter carries an unused bibliography custom-XML part.
    # This authored template has no bibliography; omit that unrelated starter metadata.
    for relationship in list(doc.part.rels.values()):
        if relationship.reltype.endswith("/customXml"):
            doc.part.drop_rel(relationship.rId)
    properties = doc.core_properties
    properties.author = "NuaForma contributors"
    properties.title = "NuaForma example template"
    properties.subject = "Project summary with nested template logic"
    properties.last_modified_by = "NuaForma contributors"
    import datetime
    properties.created = properties.modified = datetime.datetime(2026, 10, 13, tzinfo=datetime.timezone.utc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = section.footer_distance = Inches(.3)
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    existing = rpr.find(qn("w:rFonts"))
    if existing is not None:
        rpr.remove(existing)
    element(rpr, "w:rFonts", {"w:ascii": FONT, "w:hAnsi": FONT, "w:eastAsia": FONT, "w:cs": FONT})
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = FONT
            style.font.color.rgb = RGBColor(0, 0, 0)
            fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is not None:
                for key in list(fonts.attrib):
                    if "Theme" in key or "theme" in key:
                        del fonts.attrib[key]
                fonts.set(qn("w:ascii"), FONT)
                fonts.set(qn("w:hAnsi"), FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    doc.styles["Title"].font.size = Pt(27)
    doc.styles["Title"].paragraph_format.space_after = Pt(13)
    title_borders = element(doc.styles["Title"].element.get_or_add_pPr(), "w:pBdr", {})
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        element(title_borders, "w:" + edge, {"w:val": "nil"})
    for name, size in (("Heading 1", 17), ("Heading 2", 13)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(14)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    section.header.paragraphs[0].text = "NuaForma example kit"
    section.header.paragraphs[0].runs[0].font.size = Pt(9)
    section.footer.paragraphs[0].text = "Reference {reference}    •    Prepared for {customer.name}"
    section.footer.paragraphs[0].runs[0].font.size = Pt(9)
    picture = doc.add_paragraph().add_run().add_picture(str(ROOT / "assets/branding/nuaforma-avatar.jpg"), width=Inches(.65))
    picture._inline.docPr.set("descr", "NuaForma organization logo")
    doc.add_paragraph("Project delivery summary", "Title")
    doc.add_paragraph("This summary records the requested work, the team responsible and the current delivery status.")
    doc.add_paragraph("Reference {reference}  |  Issued {issuedOn | date(\"dd MMM uuuu\")}")
    doc.add_paragraph("Prepared for {trim(customer.name)} at {customer.company}.")
    doc.add_paragraph('Contact {customer.email ?? "No email supplied"}. Phone {customer.phone ?? "Not provided"}.')
    doc.add_paragraph('{#if equalsIgnoreCase(trim(status), "approved")}Approved for delivery{#else}Awaiting approval{/if}. {#if contains(description, "urgent", true)}Urgent handling requested.{#else}Standard handling.{/if}')
    doc.add_paragraph('{#if trim(customer.email ?? "") == ""}The contact email is missing, null or blank.{#else}Contact email available.{/if}')
    doc.add_paragraph("Requested work", "Heading 1")
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = (3.15, .6, 1.1, 1.15)
    table_width = Inches(sum(widths))
    table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:w"), str(table_width.twips))
    table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = Inches(width)
    for cell, title in zip(table.rows[0].cells, ("Item", "Qty", "Unit price", "Amount")):
        cell.text = title
    element(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader", {})
    def marker_row(text):
        cells = table.add_row().cells
        cell = cells[0].merge(cells[-1])
        cell.width = table_width
        cell.text = text
    marker_row("{#each lines as line}")
    cells = table.add_row().cells
    for cell, width in zip(cells, widths):
        cell.width = Inches(width)
    for cell, value in zip(cells, ('{line.description}{#if line.optional} (optional){/if}', '{line.quantity}', '{line.unitPrice | number("#,##0.00")}', '{line.amount | number("#,##0.00")}')):
        cell.text = value
    marker_row("{/each}")
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = element(props, "w:tcBorders", {})
            for edge in ("top", "left", "bottom", "right"):
                element(borders, "w:" + edge, {"w:val": "single", "w:sz": 4, "w:color": "D9D9D9"})
            margins = element(props, "w:tcMar", {})
            for edge in ("top", "left", "bottom", "right"):
                element(margins, "w:" + edge, {"w:w": 95, "w:type": "dxa"})
            if index == 0:
                element(props, "w:shd", {"w:fill": "173846", "w:val": "clear"})
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    if index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph('Total {currency} {total | number("#,##0.00")}. Amounts are supplied by the caller, not calculated in the template.')
    doc.add_paragraph("{#if showProjects}")
    doc.add_paragraph("Teams and deliverables", "Heading 1")
    doc.add_paragraph("{#each projects as project}")
    doc.add_paragraph("{project.name}", "Heading 2")
    doc.add_paragraph("{#if project.active}")
    doc.add_paragraph('Owner {project.owner.name ?? "Unassigned"}. Team {project.team}.')
    doc.add_paragraph("{#each project.tasks as task}")
    doc.add_paragraph('{loop.number}. {task.name}: {#if task.done}Complete{#else}{#if task.blocked}Blocked{#else}In progress{/if}{/if}. {#if startsWith(task.code, "NF-")}Internal reference {task.code}.{#else}External reference {task.code}.{/if}')
    doc.add_paragraph("{#if task.hasNotes}")
    doc.add_paragraph("{#each task.notes as note}")
    paragraph = doc.add_paragraph('Note {loop.number} for {task.name}: {#if contains(note, "review", true)}Review requested: {/if}{note}')
    paragraph.paragraph_format.left_indent = Inches(.2)
    doc.add_paragraph("{/each}")
    doc.add_paragraph("{/if}")
    doc.add_paragraph("{/each}")
    doc.add_paragraph("{#else}")
    doc.add_paragraph("This project is paused and its tasks are omitted.")
    doc.add_paragraph("{/if}")
    doc.add_paragraph("{/each}")
    doc.add_paragraph("{#else}")
    doc.add_paragraph("Project details are not included in this summary.")
    doc.add_paragraph("{/if}")
    doc.add_paragraph("Delivery details", "Heading 1")
    doc.add_paragraph('Destination {customer.address.city ?? "City not supplied"}, {customer.address.country ?? "Country not supplied"}.')
    doc.add_paragraph('{#if (customer.nickname ?? null) == null}No preferred name was supplied.{#else}Preferred name {customer.nickname}.{/if}')
    doc.add_paragraph('{#if endsWith(delivery.filename, ".pdf", true)}PDF delivery requested.{#else}Editable document delivery requested.{/if}')
    doc.add_paragraph('Message {message ?? "No additional message"}')
    # An original, opaque illustration keeps the same template within the PDF/A-1b image profile.
    illustration = Image.new("RGB", (1200, 200), "#eef4f1")
    draw = ImageDraw.Draw(illustration)
    for index, label in enumerate(("TEMPLATE", "STRUCTURED DATA", "DOCUMENT")):
        start = 30 + index * 395
        draw.rounded_rectangle((start, 35, start + 350, 165), radius=15, fill="#173846" if index == 2 else "#007e83")
        draw.text((start + 30, 89), label, fill="white", font_size=24)
    stream = io.BytesIO()
    illustration.save(stream, format="JPEG", quality=95)
    stream.seek(0)
    picture = doc.add_paragraph().add_run().add_picture(stream, width=table_width)
    picture._inline.docPr.set("descr", "Template and structured data produce a document")
    delivery_end = len(doc.paragraphs)
    if include_content_image:
        doc.add_paragraph('{#if (image ?? null) != null}')
        doc.add_paragraph('Supplied illustration', 'Heading 1')
        doc.add_paragraph('The image below comes from the JSON data and fits within 60 by 25 millimetres without changing its aspect ratio.')
        doc.add_paragraph('{image | image(60, 25, true)}')
        doc.add_paragraph('{/if}')
    delivery_started = False
    for paragraph in doc.paragraphs[:delivery_end - 1]:
        if paragraph.text == "Delivery details":
            delivery_started = True
        if delivery_started:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
    # Include table/latent style font declarations, not only paragraph/run styles.
    for fonts in doc.styles.element.iter(qn("w:rFonts")):
        for key in list(fonts.attrib):
            if key.lower().endswith("theme"):
                del fonts.attrib[key]
        for family in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + family), FONT)
    doc.save(OUT / "nuaforma-example-template.docx")
    normalize_zip(OUT / "nuaforma-example-template.docx")


def build_requests():
    data = {
        "reference": "NF-001", "issuedOn": "2026-10-13", "status": " APPROVED ",
        "customer": {"name": " Ada Lovelace ", "company": "Example Studio", "email": "ada@example.invalid", "phone": None, "nickname": None, "address": {"city": "Utrecht", "country": "Netherlands"}},
        "description": "Regular project delivery", "currency": "EUR", "total": 875.50,
        "lines": [{"description": "Template design", "quantity": 1, "unitPrice": 500, "amount": 500, "optional": False}, {"description": "Document delivery", "quantity": 3, "unitPrice": 125.1667, "amount": 375.50, "optional": True}],
        "showProjects": True,
        "projects": [{"name": "Customer onboarding", "active": True, "team": "Delivery", "owner": {"name": "Sam"}, "tasks": [
            {"code": "NF-T01", "name": "Prepare template", "done": True, "blocked": False, "hasNotes": True, "notes": ["Design review complete", "Images and tables retained"]},
            {"code": "EXT-02", "name": "Confirm delivery", "done": False, "blocked": True, "hasNotes": False, "notes": None}]},
            {"name": "Future rollout", "active": False, "team": "Operations", "owner": {"name": None}, "tasks": []}],
        "delivery": {"filename": "summary.pdf"}, "message": "Ready for the next review."
    }
    def envelope(values, profile="pdfa-3b", watermark=None, docx=False):
        options = {"locale": "en-GB", "timeZone": "Europe/Amsterdam"}
        if watermark is not None:
            options["watermark"] = watermark
        result = {"input": {"format": "docx"}, "output": {"format": "docx" if docx else "pdf", "options": {} if docx else {"profile": profile}}, "filename": "example", "data": values, "options": options}
        if docx:
            result["providerId"] = "nuaforma-docx"
        return result
    scenarios = []
    def scenario(file, title, profile, description, expected, request):
        (OUT / file).write_text(json.dumps(request, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
        scenarios.append(dict(file=file, title=title, profile=profile, description=description, expected=expected))
    scenario("01-final-pdfa3b.json", "Final document", "PDF/A-3b · no watermark", "Approved status, two work items and active/paused projects. Exercises nested if → each → if → each → if content.", "Two work rows; complete and blocked tasks; review note; paused-project message; no watermark.", envelope(copy.deepcopy(data)))
    draft = copy.deepcopy(data)
    draft.update(reference="NF-002", status="draft", description="URGENT customer review")
    draft["customer"]["email"] = None
    draft["customer"].pop("phone")
    draft["projects"][0]["tasks"][1].update(blocked=False, hasNotes=True, notes=["REVIEW delivery address", "Confirm the contact details"])
    watermark = {"text": '{#if equalsIgnoreCase(trim(status), "approved")}FINAL{#else}DRAFT{/if} {reference}', "color": "#00858A", "opacity": .22, "rotation": -45}
    scenario("02-draft-watermark.json", "Draft with a watermark", "PDF/A-2b · dynamic watermark", "A null email and a missing phone use fallbacks; uppercase URGENT still matches. More nested notes and an in-progress task.", "DRAFT NF-002 watermark, awaiting-approval and urgent text, safe email/phone fallbacks.", envelope(draft, "pdfa-2b", watermark))
    minimal = copy.deepcopy(data)
    minimal.update(reference="NF-003", showProjects=False, lines=[], total=0, message=None)
    minimal["customer"].update(email="   ", address=None)
    minimal["customer"].pop("nickname")
    minimal["delivery"]["filename"] = "SUMMARY.PDF"
    scenario("03-null-empty-standard.json", "Null, missing and empty", "Standard PDF 1.7 · no watermark", "Blank email, null address/message, missing nickname, empty work rows and an excluded project section. Header row remains.", "A header-only table, blank-email warning, address/message fallbacks, no project details.", envelope(minimal, "standard"))
    opaque = copy.deepcopy(data)
    opaque.update(reference="NF-004", status="draft")
    scenario("04-opaque-pdfa1b.json", "Opaque draft watermark", "PDF/A-1b · restricted profile", "The same template uses opaque JPEG images and an opaque text watermark. PDF/A-1b cannot accept transparency.", "PDF/A-1b with DRAFT NF-004 watermark at full opacity; no transparency flattening.", envelope(opaque, "pdfa-1b", dict(watermark, opacity=1)))
    editable = copy.deepcopy(data)
    editable.update(reference="NF-005", status="draft")
    editable["delivery"]["filename"] = "summary.docx"
    scenario("05-editable-word.json", "Editable Word output", "DOCX · dynamic watermark", "Choose the DOCX provider explicitly. The nested logic and table rows use the same data language.", "An editable Word document with a resolved draft watermark and retained tables/images.", envelope(editable, watermark=watermark, docx=True))
    empty = copy.deepcopy(data)
    empty.update(reference="NF-006", projects=[], message="Literal markers stay literal: {customer.name}")
    scenario("06-empty-projects-pdfa2b.json", "Empty list and literal text", "PDF/A-2b · no watermark", "An empty project list expands to no projects. A data string containing a placeholder is not evaluated a second time.", "No repeated project content; the literal text {customer.name} remains in the message.", envelope(empty, "pdfa-2b"))
    (OUT / "scenarios.json").write_text(json.dumps(scenarios, indent=2) + "\n", encoding="utf-8", newline="\n")


def package():
    files = [OUT / "nuaforma-example-template.docx", OUT / "scenarios.json", *sorted(OUT.glob("0*.json")), OUT / "EXAMPLES.md"]
    files.extend(path for name in ("bruno", "postman") for path in sorted((OUT / name).rglob("*")) if path.is_file())
    with zipfile.ZipFile(OUT / "example-kit.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for file in files:
            assert file.is_file(), file
            info = zipfile.ZipInfo(file.relative_to(OUT).as_posix(), (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, file.read_bytes())


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package-only", action="store_true", help="Reuse the existing Word/JSON examples; rebuild collections and archives only")
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    if not args.package_only:
        build_template()
        build_requests()
    build_client_collections()
    package()
    package_complete_kit()
    scenario_count = len(json.loads((OUT / "scenarios.json").read_text(encoding="utf-8")))
    print(f"Packaged {scenario_count} examples, client collections and complete starter kit")
